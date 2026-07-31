# -*- coding: utf-8 -*-
"""Eksport arkusza Internet_Koszty → JSON + TS pod widok Operacje → Finanse."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

SRC = Path(r"d:\Users\Biuro\Downloads\Internet_Koszty_v8_kaniec (2).xlsx")
OUT_PUBLIC = Path(r"C:\Users\Biuro\Projects\katalog\public\data\finance-koszty.json")
OUT_TS = Path(r"C:\Users\Biuro\Projects\katalog\src\data\financeKosztyData.ts")

CHANNEL_COLS = [
    "Allegro kenochemcom",
    "Allegro czystomania.pl",
    "Allegro sonax_sklep",
    "Kenochem.com",
    "Czystomania.pl",
    "Sonax.sklep.pl",
    "Erli",
    "Empik",
    "Kaufland",
    "Zamówienia telefoniczne/ręczne/mail",
]


def num(v):
    if v is None or v == "" or v == "-":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").replace(" ", ""))
    except ValueError:
        return 0.0


def month_key(v) -> str | None:
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m")
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    return None


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    dash = wb["Dashboard"]
    pod = wb["Podsumowanie_mies."]
    spr = wb["Sprzedaż_Baselinker"]
    rej = wb["Rejestr_kosztów"]

    selected = month_key(dash.cell(2, 2).value) or "2026-05"

    months = []
    for r in range(5, 200):
        mk = month_key(pod.cell(r, 1).value)
        if not mk:
            continue
        sprzedaz = num(pod.cell(r, 2).value)
        if sprzedaz <= 0:
            continue
        months.append(
            {
                "month": mk,
                "sprzedaz": round(sprzedaz, 2),
                "kosztTowaru": round(num(pod.cell(r, 3).value), 2),
                "marzaNetto": round(num(pod.cell(r, 4).value), 2),
                "marketplace": round(num(pod.cell(r, 5).value), 2),
                "dostawa": round(num(pod.cell(r, 6).value), 2),
                "operacyjne": round(num(pod.cell(r, 7).value), 2),
                "kosztyRazem": round(num(pod.cell(r, 8).value), 2),
                "marzaPct": num(pod.cell(r, 9).value),
                "wynikDoSprzedazy": num(pod.cell(r, 10).value),
                "kosztDoSprzedazy": num(pod.cell(r, 11).value),
                "wynikNetto": round(num(pod.cell(r, 13).value), 2),
            }
        )

    channels_by_month: dict[str, list] = {}
    for r in range(5, 200):
        mk = month_key(spr.cell(r, 1).value)
        if not mk:
            continue
        items = []
        for i, name in enumerate(CHANNEL_COLS, start=2):
            amt = num(spr.cell(r, i).value)
            if amt:
                items.append({"channel": name, "amount": round(amt, 2)})
        items.sort(key=lambda x: -x["amount"])
        if items:
            channels_by_month[mk] = items

    sources_by_month: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for r in range(5, rej.max_row + 1):
        mk = month_key(rej.cell(r, 1).value) or month_key(rej.cell(r, 2).value)
        if not mk:
            continue
        source = str(rej.cell(r, 4).value or "").strip()
        amt = num(rej.cell(r, 5).value)
        if not source or not amt:
            continue
        sources_by_month[mk][source] += amt

    sources_out = {
        mk: [
            {"name": k, "amount": round(v, 2)}
            for k, v in sorted(bag.items(), key=lambda x: -x[1])
        ]
        for mk, bag in sources_by_month.items()
    }

    if selected not in {m["month"] for m in months} and months:
        selected = months[-1]["month"]

    payload = {
        "meta": {
            "title": "Finanse sklepu Kenochem",
            "note": "Wynik netto ≈ sprzedaż − koszt towaru − koszty biznesu. Przybliżenie operacyjne z arkusza kosztów — nie pełna księgowość.",
            "currency": "PLN",
            "sourceFile": SRC.name,
            "defaultMonth": selected,
        },
        "months": months,
        "channelsByMonth": channels_by_month,
        "sourcesByMonth": sources_out,
    }

    text = json.dumps(payload, ensure_ascii=False, indent=2)
    OUT_PUBLIC.parent.mkdir(parents=True, exist_ok=True)
    OUT_PUBLIC.write_text(text, encoding="utf-8")

    ts = (
        "/* Wygenerowane przez scripts/export-finance-koszty.py — nie edytuj ręcznie. */\n"
        "import type { FinanceKosztyData } from '../lib/financeTypes';\n\n"
        f"export const financeKosztyData: FinanceKosztyData = {text} as const;\n"
    )
    # `as const` may conflict with typed annotation - use without as const
    ts = (
        "/* Wygenerowane przez scripts/export-finance-koszty.py — nie edytuj ręcznie. */\n"
        "import type { FinanceKosztyData } from '../lib/financeTypes';\n\n"
        f"export const financeKosztyData = {text} satisfies FinanceKosztyData;\n"
    )
    OUT_TS.parent.mkdir(parents=True, exist_ok=True)
    OUT_TS.write_text(ts, encoding="utf-8")
    print(f"OK months={len(months)} selected={selected} -> {OUT_PUBLIC.name}, {OUT_TS.name}")


if __name__ == "__main__":
    main()
